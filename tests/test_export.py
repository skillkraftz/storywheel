"""Export: the Shunn .docx is read back and checked against the format; the other formats; the commands."""
import datetime
import json
import os
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Inches, Pt

from storywheel import export, settings, vault

TODAY = datetime.date.today().isoformat()
ROOT = Path(__file__).resolve().parent.parent
TEXT = ("Stacie ran down the road, past the mill and the dry creek bed, and did not look back once. "
        "It was *very* dry, and **nothing** moved.")


@pytest.fixture
def story(home):
    settings.save_global({"legal_name": "Andrew T. Writer", "author_name": "Andy Writer",
                          "address": "1 Main Street\nDry Fork, WY 82000", "email": "andy@example.com", "phone": "555-0100"})
    u = vault.create_universe("Thornwood")
    s = u.new_story("The Last Clause")
    paragraphs = "\n\n".join([TEXT] * 6)
    s.add_scene("Opening", paragraphs + "\n\n* * *\n\nThe next morning the sheriff came.")
    s.add_scene("Second", "The clause was hers all along.")
    settings.save_story(s.path, {"title_keyword": "Clause"})
    return s


def read(story, **settings_changes):
    if settings_changes:
        st = settings.load_story(story.path)
        st.update(settings_changes)
        settings.save_story(story.path, st)
    result = export.export(story, "docx")
    return docx.Document(result["path"]), result


def texts(d):
    return [p.text for p in d.paragraphs]


# --- the Shunn short-story .docx, read back ----------------------------------------------------------------------

def test_page_size_and_margins(story):
    d, _ = read(story)
    s = d.sections[0]
    assert (s.page_width, s.page_height) == (Inches(8.5), Inches(11))
    assert (s.left_margin, s.right_margin, s.top_margin, s.bottom_margin) == (Inches(1),) * 4


def test_font_is_twelve_point_times_new_roman_everywhere(story):
    d, _ = read(story)
    normal = d.styles["Normal"]
    assert normal.font.name == "Times New Roman" and normal.font.size == Pt(12)
    rfonts = normal.element.rPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    assert {v for k, v in rfonts.attrib.items()} == {"Times New Roman"}
    for p in d.paragraphs:
        for r in p.runs:
            assert r.font.name in (None, "Times New Roman") and r.font.size in (None, Pt(12))


def test_courier_is_a_setting(story):
    d, _ = read(story, font="Courier New")
    assert d.styles["Normal"].font.name == "Courier New"


def test_body_is_double_spaced_with_half_inch_indents_and_no_space_between_paragraphs(story):
    d, _ = read(story)
    normal = d.styles["Normal"].paragraph_format
    assert normal.line_spacing_rule == WD_LINE_SPACING.DOUBLE and normal.space_before == Pt(0) and normal.space_after == Pt(0)
    body = [p for p in d.paragraphs if p.text.startswith("Stacie ran")]
    assert len(body) == 6
    for p in body:
        assert p.paragraph_format.first_line_indent == Inches(0.5)
        assert p.paragraph_format.line_spacing_rule in (None, WD_LINE_SPACING.DOUBLE)       # inherits double
        assert p.paragraph_format.space_after in (None, Pt(0))
        assert p.alignment in (None, WD_ALIGN_PARAGRAPH.LEFT)


def test_first_page_block_has_name_address_email_phone_and_the_word_count_at_the_right(story):
    d, result = read(story)
    ps = d.paragraphs
    assert ps[0].text.startswith("Andrew T. Writer\tabout ") and ps[0].text.endswith(" words")
    block = [p.text for p in ps[:5]]
    assert block[1:] == ["1 Main Street", "Dry Fork, WY 82000", "andy@example.com", "555-0100"]
    stops = ps[0].paragraph_format.tab_stops
    assert len(stops) == 1 and stops[0].position == Inches(6.5) and str(stops[0].alignment).startswith("RIGHT")
    for p in ps[:5]:
        assert p.paragraph_format.line_spacing_rule == WD_LINE_SPACING.SINGLE          # the block is single spaced
    words = result["words"]
    assert ps[0].text.split("\t")[1] == f"about {export.round_words(words):,} words"


def test_word_count_rounding():
    assert [export.round_words(n) for n in (3, 47, 99, 100, 149, 150, 991, 1250, 3499, 3500)] == \
           [10, 50, 100, 100, 100, 200, 1000, 1300, 3500, 3500]


def test_title_is_centered_about_halfway_down_with_the_byline_under_it(story):
    d, _ = read(story)
    ps = d.paragraphs
    i = next(n for n, p in enumerate(ps) if p.text == "The Last Clause")
    title, byline = ps[i], ps[i + 1]
    assert title.alignment == WD_ALIGN_PARAGRAPH.CENTER and byline.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert byline.text == "by Andy Writer"
    # 4.5 inches below the top margin is halfway down the page; the block above takes about 5 lines
    space = title.paragraph_format.space_before.pt
    assert 4.5 * 72 - 6 * 14 <= space + 5 * 13.8 <= 4.5 * 72 + 20
    assert i == 5                                                                       # straight after the block
    assert ps[i + 2].text.startswith("Stacie ran") or ps[i + 2].text == ""


def test_header_on_pages_two_onwards_is_surname_title_and_a_page_number_field(story):
    d, _ = read(story)
    s = d.sections[0]
    assert s.different_first_page_header_footer is True
    assert s.first_page_header.paragraphs[0].text == ""                                  # nothing on page 1
    h = s.header.paragraphs[0]
    assert h.alignment == WD_ALIGN_PARAGRAPH.RIGHT and h.text == "Writer / The Last Clause / "
    xml = h._p.xml
    assert "PAGE" in xml and 'w:fldCharType="begin"' in xml and 'w:fldCharType="end"' in xml
    assert h.paragraph_format.line_spacing_rule == WD_LINE_SPACING.SINGLE


def test_the_header_can_use_shunns_keyword_which_falls_back_to_the_first_real_word_of_the_title(story):
    st = settings.load_story(story.path)
    st["export_header"] = "keyword"
    st["title_keyword"] = "Clause"
    settings.save_story(story.path, st)
    d, _ = read(story)
    assert d.sections[0].header.paragraphs[0].text == "Writer / Clause / "
    st["title_keyword"] = ""
    settings.save_story(story.path, st)
    d, _ = read(story)
    assert d.sections[0].header.paragraphs[0].text == "Writer / Last / "


def test_a_scene_break_is_a_centered_hash_and_the_story_ends_with_a_centered_END(story):
    d, _ = read(story)
    hashes = [p for p in d.paragraphs if p.text == "#"]
    assert len(hashes) == 1 and hashes[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert hashes[0].paragraph_format.first_line_indent in (None, Inches(0))
    last = d.paragraphs[-1]
    assert last.text == "END" and last.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert "* * *" not in "\n".join(texts(d)) and "*" not in "\n".join(texts(d))


def test_italics_and_bold_stay_italics_and_bold(story):
    d, _ = read(story)
    p = next(p for p in d.paragraphs if p.text.startswith("Stacie ran"))
    by_text = {r.text: r for r in p.runs}
    assert by_text["very"].italic is True and by_text["nothing"].bold is True
    assert not by_text["very"].bold and not by_text["nothing"].italic
    assert any(r.italic is None and r.text.startswith("Stacie ran") for r in p.runs)


def test_all_the_text_is_there_in_order(story):
    d, _ = read(story)
    all_text = texts(d)
    body = [t for t in all_text[all_text.index("by Andy Writer") + 1:] if t]
    assert body[0].startswith("Stacie ran down the road") and body[-2] == "The clause was hers all along." and body[-1] == "END"
    assert "The next morning the sheriff came." in body
    assert body.index("#") < body.index("The next morning the sheriff came.") < body.index("The clause was hers all along.")


def test_the_document_properties_name_the_story(story):
    d, _ = read(story)
    assert d.core_properties.title == "The Last Clause" and d.core_properties.author == "Andy Writer"


def test_missing_author_details_are_reported_not_hidden(home):
    u = vault.create_universe("U")
    s = u.new_story("Tale")
    s.add_scene("A", "Hello world.")
    d, result = read(s)
    assert any("No author name" in w and "anonymously" in w for w in result["warnings"])
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Your Name" not in text and "Author" not in text and not any(p.text.startswith("by ") for p in d.paragraphs)
    assert d.sections[0].header.paragraphs[0].text == "Tale / "


def test_an_empty_manuscript_is_refused_plainly(home):
    u = vault.create_universe("U")
    s = u.new_story("Tale")
    with pytest.raises(export.ExportError, match="empty"):
        export.export(s, "docx")
    with pytest.raises(export.ExportError, match="Unknown format"):
        export.export(s, "rtf")


def test_a_novel_starts_each_scene_file_as_a_chapter_on_a_new_page(story):
    d, result = read(story, format="novel")
    chapters = [p for p in d.paragraphs if p.text.startswith("Chapter ")]
    assert [p.text for p in chapters] == ["Chapter 1", "Chapter 2"]
    assert chapters[1].paragraph_format.page_break_before is True
    assert any("Novel layout is partial" in w for w in result["warnings"])


def test_the_docx_opens_in_libreoffice_and_has_the_right_pages(story):
    soffice = shutil.which("soffice")
    if not soffice:
        pytest.skip("LibreOffice isn't installed")
    result = export.export(story, "pdf")
    pdf = Path(result["path"])
    assert pdf.read_bytes().startswith(b"%PDF")
    if shutil.which("pdftotext"):
        page1 = subprocess.run(["pdftotext", "-f", "1", "-l", "1", "-layout", str(pdf), "-"], capture_output=True, text=True).stdout
        assert "Andrew T. Writer" in page1 and "by Andy Writer" in page1 and "Writer / The Last Clause" not in page1
        full = subprocess.run(["pdftotext", "-layout", str(pdf), "-"], capture_output=True, text=True).stdout
        assert "Writer / The Last Clause / 2" in full and "END" in full and "#" in full


# --- the other formats ----------------------------------------------------------------------------------------------

def test_markdown_compiles_the_scenes_with_a_title_block(story):
    r = export.export(story, "md")
    text = Path(r["path"]).read_text()
    assert text.startswith("# The Last Clause\n\n*by Andy Writer*\n\n") and "*very*" in text and "* * *" in text
    assert text.index("The next morning") < text.index("The clause was hers")


def test_plain_text_has_no_markup_at_all(story):
    r = export.export(story, "txt")
    text = Path(r["path"]).read_text()
    assert "*" not in text and "#" in text.split("\n\n") and text.rstrip().endswith("END")
    assert "very dry" in text and text.startswith("The Last Clause\nby Andy Writer\n")


def test_plain_text_for_the_clipboard(story):
    text = export.plain_text(story)
    assert "*" not in text and "\n\n#\n\n" in text and text.startswith("Stacie ran")


def test_odt_is_made_with_libreoffice(story):
    if not shutil.which("soffice"):
        pytest.skip("LibreOffice isn't installed")
    r = export.export(story, "odt")
    import zipfile
    z = zipfile.ZipFile(r["path"])
    assert z.read("mimetype") == b"application/vnd.oasis.opendocument.text"


def test_without_libreoffice_odt_and_pdf_say_so_and_keep_the_docx(story, monkeypatch):
    monkeypatch.setattr(export.shutil, "which", lambda name: None)
    with pytest.raises(export.ExportError, match="LibreOffice"):
        export.export(story, "pdf")
    assert (Path(os.environ["STORYWHEEL_MANUSCRIPTS"]) / "The Last Clause" / f"The Last Clause {TODAY}.docx").exists()


def test_a_screenplay_is_a_marked_stub(story):
    st = settings.load_story(story.path)
    st["format"] = "screenplay"
    settings.save_story(story.path, st)
    r = export.export(story, "docx")
    assert r["format"] == "fountain" and r["path"].endswith(".fountain")
    assert any("stub" in w for w in r["warnings"])
    assert Path(r["path"]).read_text().startswith("Title: The Last Clause\nAuthor: Andy Writer\n")


# --- commands ---------------------------------------------------------------------------------------------------------------

def cli(args, home):
    env = dict(os.environ, STORYWHEEL_HOME=str(home / "home"), STORYWHEEL_LIBRARY=str(home / "library"),
               STORYWHEEL_OUT=str(home / "out"), PYTHONIOENCODING="utf-8")
    return subprocess.run([sys.executable, "-m", "storywheel", *args], capture_output=True, text=True, env=env,
                          cwd=ROOT, encoding="utf-8")


def test_manuscript_export_and_text_commands(home, story):
    res = cli(["manuscript", "export", "thornwood/the-last-clause", "--format", "docx", "--json"], home)
    data = json.loads(res.stdout)
    assert data["format"] == "docx" and data["path"].endswith(f"The Last Clause/The Last Clause {TODAY}.docx") and Path(data["path"]).exists()
    out = cli(["manuscript", "export", "thornwood/the-last-clause", "--format", "md", "--out", str(home / "elsewhere")], home)
    assert "Wrote" in out.stdout and (home / "elsewhere" / f"The Last Clause {TODAY}.md").exists()
    plain = cli(["manuscript", "text", "thornwood/the-last-clause"], home).stdout
    assert plain.startswith("Stacie ran") and "*" not in plain
    bad = cli(["manuscript", "export", "thornwood/the-last-clause", "--format", "rtf", "--json"], home)
    assert bad.returncode == 1 and "Unknown format" in json.loads(bad.stdout)["error"]


# --- from the Builder and the Writer -------------------------------------------------------------------------------------------

def test_the_builder_exports_with_a_chosen_format_and_shows_where(home, story):
    import asyncio
    from storywheel import builder, fill

    async def go():
        app = builder.BuilderApp(engine_factory=lambda u: fill.make_engine(u, seed=1), universe="thornwood", story="the-last-clause")
        async with app.run_test(size=(220, 55)) as pilot:
            await pilot.pause()
            await pilot.press("x")
            await pilot.pause()
            asked = type(app.screen).__name__
            await pilot.press("enter")                      # the first choice: docx
            await pilot.pause()
            from conftest import screen_text
            return asked, " ".join(screen_text(app).split()), app.screen_ref.last_export
    asked, text, result = asyncio.run(go())
    assert asked == "ChoiceScreen" and "Exported" in text and f"The Last Clause {TODAY}.docx" in text
    assert Path(result["path"]).exists()


def test_the_builder_says_so_when_there_is_nothing_to_export(home):
    import asyncio
    from storywheel import builder
    u = vault.create_universe("U")
    u.new_story("Empty")

    async def go():
        app = builder.BuilderApp(universe="u", story="empty")
        async with app.run_test(size=(220, 55)) as pilot:
            await pilot.pause()
            await pilot.press("x", "enter")
            await pilot.pause()
            from conftest import screen_text
            return " ".join(screen_text(app).split())
    assert "manuscript is empty" in asyncio.run(go())


def test_the_builder_copies_the_manuscript_as_plain_text(home, story, monkeypatch):
    import asyncio
    from storywheel import builder, clipboard
    got = []
    monkeypatch.setattr(clipboard, "copy", lambda text, app=None: got.append(text) or "wl-copy")
    async def go():
        app = builder.BuilderApp(universe="thornwood", story="the-last-clause")
        async with app.run_test(size=(220, 55)) as pilot:
            await pilot.pause()
            await pilot.press("C")
            await pilot.pause()
    asyncio.run(go())
    assert got and got[0] == export.plain_text(story)


def test_export_from_inside_the_writer_calls_the_same_code(home, story):
    from test_writer import run_lua
    if shutil.which("nvim") is None:
        pytest.skip("Neovim isn't installed")
    r = run_lua(story, 'local d = require("sw").export("docx"); R.path = d and d.path; R.fmt = d and d.format')
    assert r["fmt"] == "docx" and Path(r["path"]).exists()
    d = docx.Document(r["path"])
    assert d.paragraphs[0].text.startswith("Andrew T. Writer")
