"""Batch 20: every Writer action has a notepad key, and Alt+C / Ctrl+E center lines (stored >text<, shown centered, exported centered)."""
import pytest

from storywheel import export, settings, vault, writer
from test_notepad import AT, LINES, run

pytestmark = pytest.mark.skipif(writer.check() is not None, reason="Neovim 0.10+ is not installed")


@pytest.fixture
def story(home):
    s = vault.create_universe("Keys", ["noir"]).new_story("Tale", {"format": "short-story"})
    s.append_scene("Opening", "First line.\nSecond line.\nThird line.")
    return s


def text(lines):
    return "vim.api.nvim_buf_set_lines(0, 0, -1, false, { %s })\n" % ", ".join('"%s"' % l for l in lines)


def test_the_notepad_keys_run_their_actions_while_typing(story):
    r = run(story, "", "<A-t>" + "<A-v>", "R.tw = require('sw.prose').typewriter; R.inv = require('sw.prose').invisibles")
    assert r["tw"] is True and r["inv"] is True
    r = run(story, "", "<A-l>", "R.spell = vim.wo[require('sw.layout').main].spell")
    assert r["spell"] is False                                              # (it was on by default)
    r = run(story, "", "<A-n>Next<CR>", "R.n = #require('sw.story').scene_list()")      # (it asks for a title, then adds the scene)
    assert r["n"] >= 2


def test_every_action_key_is_mapped_in_insert_mode(story):
    r = run(story, "", "", """
        R.missing = {}
        local defaults = require('sw.keys_default')
        for name, lhs in pairs(defaults) do
          local m = vim.fn.maparg(lhs, 'i', false, true)
          if m == nil or next(m) == nil then R.missing[#R.missing + 1] = name end
        end
    """)
    assert r["missing"] == [] or r["missing"] == {}


def test_alt_c_centers_a_line_and_toggles_it_back(story):
    r = run(story, text(["Hello world"]) + AT % (1, 2), "<A-c>", LINES)
    assert r["lines"] == [">Hello world<"]
    r = run(story, text(["Hello world"]) + AT % (1, 2), "<A-c><C-e>", LINES)
    assert r["lines"] == ["Hello world"]


def test_ctrl_e_centers_selected_lines_only_not_scene_markers_or_blanks(story):
    r = run(story, text(["One", "***", "Two"]) + AT % (1, 0), "<C-e>", LINES)
    assert r["lines"] == [">One<", "***", "Two"]
    r = run(story, text(["One", "Two", "Three"]) + AT % (1, 0), "<C-a><C-e>", LINES)
    assert r["lines"] == [">One<", ">Two<", ">Three<"]


def test_a_centered_line_shows_centered_with_its_markers_hidden(story):
    r = run(story, text([">Hello<"]), "", """
        require('sw.prose').decorate()
        local marks = vim.api.nvim_buf_get_extmarks(0, require('sw.prose').ns, 0, -1, { details = true })
        R.pad, R.concealed = 0, 0
        for _, m in ipairs(marks) do
          if m[4].virt_text and m[4].virt_text[1][1]:match('^ +$') then R.pad = #m[4].virt_text[1][1] end
          if m[4].conceal == '' then R.concealed = R.concealed + 1 end
        end
    """)
    assert r["pad"] > 20 and r["concealed"] == 2


def test_centered_lines_are_not_extra_words_and_the_exports_center_them(story):
    story.manuscript_dir.mkdir(exist_ok=True)
    (story.manuscript_dir / "manuscript.md").write_text("* * *\n\n>The Title Page<\n\nA plain line of prose.\n", encoding="utf-8")
    assert vault.count_words(">The Title Page<\nA plain line of prose.") == 8 and vault.centered_text(">x<") == "x"
    settings.save_global({"author_name": "Andy Example"})
    import docx
    r = export.export(story, "docx")
    d = docx.Document(r["path"])
    centered = [p for p in d.paragraphs if p.text == "The Title Page"]
    assert centered and centered[0].alignment == 1 and ">" not in "".join(p.text for p in d.paragraphs)
    txt = open(export.export(story, "txt")["path"], encoding="utf-8").read()
    assert "The Title Page" in txt and ">The" not in txt and txt.split("The Title Page")[0].endswith(" ")
    md = open(export.export(story, "md")["path"], encoding="utf-8").read()
    assert '<div align="center">The Title Page</div>' in md
