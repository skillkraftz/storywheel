"""The manuscript export: one line is one paragraph, every paragraph is indented, and the new options (title, header,
anonymous, one space) do what the settings say. Read back with python-docx."""
import datetime
import os
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from storywheel import export, settings, vault


def make(text, title="The Last Clause", save=None, name="Andy Writer"):
    u = vault.get_universe("thornwood") or vault.create_universe("Thornwood")
    s = u.new_story(title)
    s.manuscript_dir.mkdir(parents=True, exist_ok=True)
    (s.manuscript_dir / "manuscript.md").write_text(text, encoding="utf-8")
    g = {"legal_name": name, "author_name": "", "address": "1 Main St\nDry Fork", "email": "a@b.c"} if name else {}
    settings.save_global(dict(settings.load_global(), **g)) if name else None
    if save:
        settings.save_story(s.path, save)
    return s


def read(s, anonymous=None, fmt="docx"):
    r = export.export(s, fmt, anonymous=anonymous)
    return docx.Document(r["path"]), r


def body(d):
    """The story's paragraphs: everything after the title (and the byline, if there is one), without END."""
    ps = d.paragraphs
    i = next(k for k, p in enumerate(ps) if p.text == d.core_properties.title) + 1
    if i < len(ps) and ps[i].text.startswith("by "):
        i += 1
    return [p for p in ps[i:] if p.text != "END"]


PASTED = "\n".join(f"Paragraph {i} of the pasted story, with some words in it for the count." for i in range(1, 51))


def test_a_pasted_story_with_one_paragraph_per_line_and_no_blank_lines_exports_as_that_many_paragraphs(home):
    s = make(PASTED)
    d, r = read(s)
    paras = body(d)
    assert len(paras) == 50
    assert [p.text for p in paras][:2] == ["Paragraph 1 of the pasted story, with some words in it for the count.",
                                          "Paragraph 2 of the pasted story, with some words in it for the count."]
    assert all(p.paragraph_format.first_line_indent == Inches(0.5) for p in paras)


def test_every_paragraph_is_indented_including_the_first_after_a_scene_break(home):
    s = make("One.\nTwo.\n***\nThree.\n* * *\nFour.\n#\nFive.\n")
    d, _ = read(s)
    paras = body(d)
    texts = [p.text for p in paras]
    assert texts == ["One.", "Two.", "#", "Three.", "#", "Four.", "#", "Five."]
    for p in paras:
        if p.text == "#":
            assert p.alignment == WD_ALIGN_PARAGRAPH.CENTER and p.paragraph_format.first_line_indent in (None, Inches(0))
        else:
            assert p.paragraph_format.first_line_indent == Inches(0.5)


def test_blank_lines_are_optional_and_mean_nothing(home):
    a = make("One.\n\nTwo.\n\n\n***\n\nThree.\n", title="A")
    b = make("One.\nTwo.\n***\nThree.\n", title="B")
    assert [p.text for p in body(read(a)[0])] == [p.text for p in body(read(b)[0])]


def test_italic_and_bold_survive_and_stars_in_text_are_not_breaks(home):
    s = make("She said *no* and **meant** it.\n***Whole line***\n***\nAfter.\n")
    d, _ = read(s)
    paras = body(d)
    assert [r.italic for r in paras[0].runs if r.text == "no"] == [True] and [r.bold for r in paras[0].runs if r.text == "meant"] == [True]
    assert [p.text for p in paras].count("#") == 1


def test_markdown_and_text_exports_keep_one_paragraph_per_line(home):
    s = make("One.\nTwo.\n***\nThree.\n")
    md = Path(export.export(s, "md")["path"]).read_text()
    assert "\n\nOne.\n\nTwo.\n\n* * *\n\nThree.\n" in md
    txt = Path(export.export(s, "txt")["path"]).read_text()
    assert "One.\n\nTwo.\n\n#\n\nThree." in txt


# --- title, header ---------------------------------------------------------------------------------------------------------------

def title_paragraph(d):
    return next(p for p in d.paragraphs if p.text == "The Last Clause")


def test_the_title_is_bold_by_default_and_plain_when_turned_off(home):
    s = make("Text.\n")
    assert all(r.bold for r in title_paragraph(read(s)[0]).runs)
    settings.save_story(s.path, {"export_title_bold": False})
    assert not any(r.bold for r in title_paragraph(read(s)[0]).runs)


def test_the_header_shows_the_full_title_or_the_keyword(home):
    s = make("Text.\n")
    assert read(s)[0].sections[0].header.paragraphs[0].text == "Writer / The Last Clause / "
    settings.save_story(s.path, {"export_header": "keyword"})
    assert read(s)[0].sections[0].header.paragraphs[0].text == "Writer / Last / "


# --- anonymous ------------------------------------------------------------------------------------------------------------------------

def test_the_anonymous_export_has_no_name_contact_byline_or_surname(home):
    s = make("Text here.\n")
    d, r = read(s, anonymous=True)
    text = "\n".join(p.text for p in d.paragraphs)
    for needle in ("Andy", "Writer", "1 Main St", "a@b.c", "by "):
        assert needle not in text.replace("Writer / ", "") or needle == "Writer" and False
    assert "Andy Writer" not in text and "1 Main St" not in text and "a@b.c" not in text and not any(p.text.startswith("by ") for p in d.paragraphs)
    assert d.sections[0].header.paragraphs[0].text == "The Last Clause / "
    assert "about 10 words" in d.paragraphs[0].text
    assert not r["warnings"] or all("anonymous" not in w for w in r["warnings"])
    assert d.core_properties.author in ("", None) or "Andy" not in d.core_properties.author


def test_the_setting_makes_every_export_anonymous(home):
    s = make("Text here.\n", save={"export_anonymous": True})
    d, _ = read(s)
    assert d.sections[0].header.paragraphs[0].text == "The Last Clause / " and "Andy" not in "\n".join(p.text for p in d.paragraphs)


def test_missing_author_details_export_anonymously_and_say_so(home):
    u = vault.create_universe("U")
    s = u.new_story("Tale")
    s.append_scene("", "Hello there.")
    d, r = read(s)
    text = "\n".join(p.text for p in d.paragraphs)
    assert "Your Name" not in text and "Author" not in text
    assert any("anonymously" in w and "Settings (F4) > You" in w for w in r["warnings"])
    assert d.sections[0].header.paragraphs[0].text == "Tale / "


def test_anonymous_md_and_txt_have_no_byline(home):
    s = make("Text.\n")
    assert "by " not in Path(export.export(s, "txt", anonymous=True)["path"]).read_text()
    assert "*by" not in Path(export.export(s, "md", anonymous=True)["path"]).read_text()


def test_the_cli_and_the_builder_have_an_anonymous_choice(home):
    import subprocess, sys
    s = make("Text here.\n")
    root = Path(__file__).resolve().parent.parent
    out = subprocess.run([sys.executable, "-m", "storywheel", "manuscript", "export", "thornwood/the-last-clause", "--anonymous", "--json"],
                         capture_output=True, text=True, env=dict(os.environ, PYTHONPATH=str(root)))
    assert out.returncode == 0, out.stderr
    import json
    d = docx.Document(json.loads(out.stdout)["path"])
    assert d.sections[0].header.paragraphs[0].text == "The Last Clause / "


# --- one space after periods ------------------------------------------------------------------------------------------------------------

def test_one_space_after_periods_is_optional(home):
    text = "He left.  She stayed!  Why?  Nobody knew.  \"Quoted.\"  Then 3.5  apples.\n"
    s = make(text)
    assert read(s)[0].paragraphs[-2].text == "He left.  She stayed!  Why?  Nobody knew.  \u201cQuoted.\u201d  Then 3.5  apples."
    settings.save_story(s.path, {"export_one_space": True})
    d, _ = read(s)
    assert d.paragraphs[-2].text == "He left. She stayed! Why? Nobody knew. \u201cQuoted.\u201d Then 3.5  apples."
    assert "He left. She stayed!" in Path(export.export(s, "txt")["path"]).read_text()
